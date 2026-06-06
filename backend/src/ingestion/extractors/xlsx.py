"""XLSX extractor — openpyxl, every sheet, computed values.

Each worksheet becomes a ``## <sheet name>`` section holding one Markdown
table (first non-empty row = header). ``data_only=True`` reads cached
formula RESULTS, not formula strings.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from src.ingestion.cleaning import clean_extracted_text
from src.ingestion.markdown import rows_to_markdown_table
from src.ingestion.types import ExtractedDocument


class XlsxExtractor:
    def extract(self, path: Path) -> ExtractedDocument:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            sections: list[str] = []
            for sheet in workbook.worksheets:
                rows = [
                    ["" if cell is None else str(cell) for cell in row]
                    for row in sheet.iter_rows(values_only=True)
                ]
                rows = [row for row in rows if any(cell.strip() for cell in row)]
                if not rows:
                    continue
                sections.append(f"## {sheet.title}\n\n{rows_to_markdown_table(rows)}")
        finally:
            workbook.close()

        markdown = clean_extracted_text("\n\n".join(sections))
        return ExtractedDocument(
            markdown=markdown,
            status="active",
            metadata={"extractor": "xlsx", "sheet_count": len(sections)},
        )
